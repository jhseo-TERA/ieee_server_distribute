# -*- coding: utf-8 -*-
"""Cross-check and optionally import the author-curated SerDes Google Sheet.

The Google Sheet remains read-only. Download/export it to an XLSX snapshot,
run this script without ``--apply`` first, and inspect the JSON report. When
``--apply`` is supplied, only high-confidence paper matches are stored as a
separate ``user_sheet`` measurement source; existing abstract/WLink rows are
never overwritten or deleted.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
import math
from pathlib import Path
import re
import sys
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from serdes_metrics import extract_performance, normalize_title, plain_text, title_similarity  # noqa: E402
from serdes_fom import regenerate_fom  # noqa: E402
from scripts.serdes_data_pipeline import (  # noqa: E402
    component_scope_for,
    db_connect,
    ensure_schema,
    link_class_for,
    modulation_for,
    sha256_text,
    upsert_evidence,
    upsert_implementation,
    upsert_measurement,
)


PRIVATE_CONFIG_PATH = ROOT / "config" / "private_sources.json"
PRIVATE_CONFIG = json.loads(PRIVATE_CONFIG_PATH.read_text(encoding="utf-8")) if PRIVATE_CONFIG_PATH.exists() else {}
SOURCE_SPREADSHEET_ID = PRIVATE_CONFIG.get("survey_spreadsheet_id", "")
SOURCE_SPREADSHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    f"{SOURCE_SPREADSHEET_ID}/edit"
) if SOURCE_SPREADSHEET_ID else ""
SYNC_VERSION = PRIVATE_CONFIG.get("survey_sync_version", "user-survey-sync-1.0")
SURVEY_KEY_PREFIX = PRIVATE_CONFIG.get("survey_key_prefix", "survey")
SURVEY_WORKBOOK_NAME = PRIVATE_CONFIG.get("survey_workbook_name", "user_survey.xlsx")
SURVEY_REPORT_NAME = PRIVATE_CONFIG.get("survey_report_name", "user_cross_validation.json")

SHEET_GIDS = {
    "SurveyData - ISSCC": 21647953,
    "SurveyData - CICC": 360600820,
    "SurveyData - VLSI": 1973591998,
    "SurveyData - ASSCC": 2040783445,
    "SurveyData - ESSERC": 809444,
    "SurveyData - RFIC": 174312969,
    "SurveyData - JSSC": 1392842880,
    "SurveyData - TCAS I": 669205916,
    "SurveyData - TCAS II": 1755992319,
}

VENUE_ALIASES = {
    "ISSCC": ("ISSCC",),
    "CICC": ("CICC",),
    "VLSI": ("VLSI-CIRCUITS", "VLSI-TECH"),
    "ASSCC": ("ASSCC",),
    "A-SSCC": ("ASSCC",),
    "ESSERC": ("ESSCIRC",),
    "ESSCIRC": ("ESSCIRC",),
    "RFIC": ("RFIC",),
    "JSSC": ("JSSC",),
    "TCAS I": ("TCAS-I",),
    "TCAS-I": ("TCAS-I",),
    "TCAS II": ("TCAS-II",),
    "TCAS-II": ("TCAS-II",),
}

REFERENCE_FIELDS = (
    "Year", "Publication", "Title", "First Author", "Affiliation", "Process",
    "Loss (dB)", "Loss Frequency (GHz)", "Speed (Gb/s)", "Power(mW)",
    "Energy/Bit [pJ]", "Main Category", "Sub Category", "Status",
)

SOURCE_PRIORITY = {
    "manual": 60,
    "user_sheet": 55,
    "pdf": 50,
    "reference_xlsx": 40,
    "abstract": 30,
    "title": 10,
}
REVIEW_PRIORITY = {
    "verified": 4,
    "reviewed": 3,
    "curated_reference": 2,
    "extracted": 1,
    "needs_review": 0,
}
COMPARE_TOLERANCES = {
    "process_nm": (0.5, 0.02),
    "lane_rate_gbps": (0.1, 0.02),
    "reported_rate_gbps": (0.1, 0.02),
    "power_mw": (0.05, 0.05),
    "energy_pj_bit": (0.01, 0.05),
    "channel_loss_db": (0.5, 0.05),
    "loss_frequency_ghz": (0.1, 0.05),
}


@dataclass(frozen=True)
class ParsedNumber:
    value: float | None
    raw: str
    qualifier: str = ""
    uncertain: bool = False


def json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def positive_number(value: Any) -> ParsedNumber:
    """Parse a positive number while retaining scope/uncertainty annotations."""
    if value is None or isinstance(value, bool):
        return ParsedNumber(None, "")
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number) or number <= 0:
            return ParsedNumber(None, str(value), "non-positive", number == 0)
        return ParsedNumber(number, str(value))
    raw = plain_text(str(value)).strip()
    if not raw or raw.lower() in {"-", "n/a", "na", "nan"}:
        return ParsedNumber(None, raw)
    match = re.match(
        r"^\s*(?P<comparator>[<>~≤≥]?)\s*(?P<number>\d+(?:\.\d+)?)"
        r"\s*(?P<qualifier>.*)$",
        raw,
    )
    if not match:
        return ParsedNumber(None, raw, raw, True)
    qualifier = plain_text(match.group("qualifier")).strip()
    comparator = match.group("comparator")
    uncertain = bool(comparator or "?" in qualifier)
    if comparator:
        return ParsedNumber(None, raw, f"{comparator}{qualifier}".strip(), True)
    number = float(match.group("number"))
    if number <= 0:
        return ParsedNumber(None, raw, qualifier or "non-positive", True)
    return ParsedNumber(number, raw, qualifier, uncertain)


def process_value(value: Any) -> tuple[str | None, float | None, str]:
    """Normalize nm/um/angstrom process entries without discarding raw text."""
    text = plain_text(str(value)).strip() if value is not None else ""
    if not text or text.lower() in {"-", "n/a", "na", "nan"}:
        return None, None, ""
    if re.fullmatch(r"\d+z", text, re.IGNORECASE):
        return text, None, "generation name; no numeric nm conversion"
    numbers = [float(item) for item in re.findall(r"\d+(?:\.\d+)?", text)]
    if not numbers:
        return text, None, "unparsed process"
    if re.fullmatch(r"\s*\d+(?:\.\d+)?\s*(?:A|Å)\s*", text, re.IGNORECASE):
        return text, numbers[0] / 10.0, "angstrom converted to nm"
    converted = [item * 1000.0 if 0.01 <= item < 1 else item for item in numbers]
    plausible = [item for item in converted if 1 <= item <= 500]
    if not plausible:
        return text, None, "no plausible nm value"
    qualifier = "mixed process; smallest node used" if len(plausible) > 1 else ""
    if any(item < 1 for item in numbers):
        qualifier = (qualifier + "; " if qualifier else "") + "um converted to nm"
    return text, min(plausible), qualifier


def fom_consistency(speed: float | None, power: float | None, energy: float | None) -> dict:
    if not speed or not power or not energy:
        return {"status": "not_checkable", "computed_pj_bit": None}
    computed = power / speed
    delta = abs(computed - energy)
    tolerance = max(0.01, 0.02 * abs(energy))
    return {
        "status": "consistent" if delta <= tolerance else "conflict",
        "computed_pj_bit": computed,
        "absolute_delta": delta,
        "relative_delta": delta / max(abs(energy), 1e-12),
        "tolerance": tolerance,
    }


def workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_reference_rows(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records: list[dict] = []
    try:
        for sheet_name, gid in SHEET_GIDS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            row_iter = sheet.iter_rows(values_only=True)
            raw_headers = next(row_iter, ())
            headers = [plain_text(value) for value in raw_headers[: len(REFERENCE_FIELDS)]]
            for row_number, values in enumerate(row_iter, start=2):
                raw_record = dict(zip(headers, values[: len(headers)]))
                title = plain_text(raw_record.get("Title"))
                year_parsed = positive_number(raw_record.get("Year"))
                if not title or year_parsed.value is None:
                    continue
                publication = plain_text(raw_record.get("Publication")) or sheet_name.rsplit("-", 1)[-1].strip()
                speed = positive_number(raw_record.get("Speed (Gb/s)"))
                power = positive_number(raw_record.get("Power(mW)"))
                energy = positive_number(raw_record.get("Energy/Bit [pJ]"))
                loss = positive_number(raw_record.get("Loss (dB)"))
                loss_frequency = positive_number(raw_record.get("Loss Frequency (GHz)"))
                process_text, process_nm, process_qualifier = process_value(raw_record.get("Process"))
                normalized_title = normalize_title(title)
                records.append({
                    "sheet": sheet_name,
                    "gid": gid,
                    "row": row_number,
                    "source_url": f"{SOURCE_SPREADSHEET_URL}?gid={gid}#gid={gid}" if SOURCE_SPREADSHEET_URL else "",
                    "year": int(year_parsed.value),
                    "publication": publication,
                    "title": title,
                    "normalized_title": normalized_title,
                    "first_author": plain_text(raw_record.get("First Author")),
                    "affiliation": plain_text(raw_record.get("Affiliation")),
                    "main_category": plain_text(raw_record.get("Main Category")),
                    "sub_category": plain_text(raw_record.get("Sub Category")),
                    "status": plain_text(raw_record.get("Status")),
                    "process_text": process_text,
                    "process_nm": process_nm,
                    "process_qualifier": process_qualifier,
                    "speed": speed,
                    "power": power,
                    "energy": energy,
                    "loss": loss,
                    "loss_frequency": loss_frequency,
                    "fom_check": fom_consistency(speed.value, power.value, energy.value),
                    "uncertain_fields": {
                        name: parsed.raw
                        for name, parsed in {
                            "speed_gbps": speed, "power_mw": power,
                            "energy_pj_bit": energy, "channel_loss_db": loss,
                            "loss_frequency_ghz": loss_frequency,
                        }.items()
                        if parsed.uncertain
                    },
                    "raw_record": {key: json_value(value) for key, value in raw_record.items() if key},
                })
    finally:
        workbook.close()
    duplicate_counts = Counter((row["sheet"], row["normalized_title"]) for row in records)
    for row in records:
        row["sheet_duplicate"] = duplicate_counts[(row["sheet"], row["normalized_title"])] > 1
    return records


def venue_aliases(publication: str, sheet_name: str = "") -> tuple[str, ...]:
    venue = plain_text(publication).upper().replace("–", "-").strip()
    if venue in VENUE_ALIASES:
        return VENUE_ALIASES[venue]
    suffix = sheet_name.rsplit("-", 1)[-1].strip().upper() if sheet_name else ""
    return VENUE_ALIASES.get(suffix, (venue,))


def fetch_candidates(cur, records: list[dict]) -> tuple[dict, dict]:
    sources = sorted({source for row in records for source in venue_aliases(row["publication"], row["sheet"])})
    years = [row["year"] for row in records]
    placeholders = ",".join(["%s"] * len(sources))
    cur.execute(
        f"""
        SELECT id, article_number, title, authors, year, source_name, url
        FROM papers
        WHERE source_system='ieee'
          AND CAST(year AS UNSIGNED) BETWEEN %s AND %s
          AND UPPER(source_name) IN ({placeholders})
        """,
        (min(years) - 1, max(years) + 1, *sources),
    )
    by_year_source: dict[tuple[int, str], list[dict]] = defaultdict(list)
    by_source_title: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for paper in cur.fetchall():
        try:
            year = int(paper.get("year") or 0)
        except (TypeError, ValueError):
            continue
        source = plain_text(paper.get("source_name")).upper()
        paper["normalized_title"] = normalize_title(paper.get("title"))
        by_year_source[(year, source)].append(paper)
        by_source_title[(source, paper["normalized_title"])].append(paper)
    return by_year_source, by_source_title


def match_record(record: dict, by_year_source: dict, by_source_title: dict,
                 threshold: float = 0.94, margin: float = 0.025) -> dict:
    aliases = venue_aliases(record["publication"], record["sheet"])
    pool = [paper for alias in aliases for paper in by_year_source.get((record["year"], alias), [])]
    exact = [paper for paper in pool if paper["normalized_title"] == record["normalized_title"]]
    if len(exact) == 1:
        return {"paper": exact[0], "method": "exact_title_year_venue", "score": 1.0, "margin": 1.0}
    alias_exact = [
        paper for alias in aliases
        for paper in by_source_title.get((alias, record["normalized_title"]), [])
        if abs(int(paper.get("year") or 0) - record["year"]) <= 1
    ]
    unique_alias_exact = {paper["id"]: paper for paper in alias_exact}
    if len(unique_alias_exact) == 1:
        paper = next(iter(unique_alias_exact.values()))
        return {"paper": paper, "method": "exact_title_venue_year_offset", "score": 0.99, "margin": 1.0}
    scored = sorted(
        ((title_similarity(record["title"], paper.get("title")), paper) for paper in pool),
        key=lambda item: item[0], reverse=True,
    )
    best_score, best = scored[0] if scored else (0.0, None)
    second_score = scored[1][0] if len(scored) > 1 else 0.0
    score_margin = best_score - second_score
    if best is not None and best_score >= threshold and score_margin >= margin:
        return {"paper": best, "method": "fuzzy_title_year_venue", "score": best_score, "margin": score_margin}
    reason = "ambiguous" if best is not None and best_score >= threshold else "unmatched"
    return {"paper": None, "method": reason, "score": best_score, "margin": score_margin,
            "best_candidate": best}


def fetch_screenings(cur) -> dict[int, dict]:
    cur.execute(
        """
        SELECT s.paper_id, s.relevance_class, s.include_in_survey, s.review_status
        FROM serdes_paper_screenings s
        """
    )
    return {int(row["paper_id"]): row for row in cur.fetchall()}


def measurement_rank(row: dict) -> tuple:
    populated = sum(row.get(field) is not None for field in COMPARE_TOLERANCES)
    return (
        REVIEW_PRIORITY.get(row.get("review_status"), 0),
        SOURCE_PRIORITY.get(row.get("source_kind"), 0),
        populated, float(row.get("overall_confidence") or 0), int(row.get("id") or 0),
    )


def fetch_existing_measurements(cur, paper_ids: list[int]) -> dict[int, list[dict]]:
    if not paper_ids:
        return {}
    placeholders = ",".join(["%s"] * len(paper_ids))
    cur.execute(
        f"""
        SELECT ip.paper_id, m.*
        FROM serdes_implementation_papers ip
        JOIN serdes_measurements m ON m.implementation_id=ip.implementation_id
        WHERE ip.paper_id IN ({placeholders}) AND m.source_kind <> 'user_sheet'
        """, tuple(paper_ids),
    )
    grouped: dict[int, list[dict]] = defaultdict(list)
    for row in cur.fetchall():
        grouped[int(row["paper_id"])].append(row)
    return grouped


def reference_metrics(record: dict, omit_fom_triple: bool = False) -> dict:
    values = {
        "process_nm": record["process_nm"],
        "reported_rate_gbps": record["speed"].value,
        "lane_rate_gbps": record["speed"].value,
        "power_mw": record["power"].value,
        "energy_pj_bit": record["energy"].value,
        "channel_loss_db": record["loss"].value,
        "loss_frequency_ghz": record["loss_frequency"].value,
    }
    if record["loss_frequency"].uncertain:
        values["loss_frequency_ghz"] = None
    if omit_fom_triple:
        for field in ("reported_rate_gbps", "lane_rate_gbps", "power_mw", "energy_pj_bit"):
            values[field] = None
    return values


def best_existing_for_field(rows: list[dict], field: str) -> dict | None:
    candidates = [row for row in rows if row.get(field) is not None]
    return max(candidates, key=measurement_rank) if candidates else None


def compare_field(new_value: float, existing: dict | None, field: str) -> dict:
    if existing is None:
        return {"status": "local_missing", "sheet": new_value, "local": None}
    local_value = float(existing[field])
    absolute_delta = abs(float(new_value) - local_value)
    absolute_tolerance, relative_tolerance = COMPARE_TOLERANCES[field]
    relative_delta = absolute_delta / max(abs(float(new_value)), abs(local_value), 1e-12)
    status = "agree" if absolute_delta <= absolute_tolerance or relative_delta <= relative_tolerance else "conflict"
    return {
        "status": status, "sheet": float(new_value), "local": local_value,
        "absolute_delta": absolute_delta, "relative_delta": relative_delta,
        "local_source": existing.get("source_kind"),
        "local_review_status": existing.get("review_status"),
        "local_measurement_id": existing.get("id"),
    }


def operating_point_key(record: dict, suffix: str = "") -> str:
    venue = re.sub(r"[^a-z0-9]+", "", record["publication"].lower())[:10] or "survey"
    stable = sha256_text(f"{record['sheet']}|{record['year']}|{record['normalized_title']}")[:16]
    return f"{SURVEY_KEY_PREFIX}-{venue}-{stable}{suffix}"[:64]


def store_record_measurement(cur, record: dict, paper: dict, workbook_hash: str,
                             *, conflict_candidate: bool = False) -> int | None:
    metrics = reference_metrics(
        record,
        omit_fom_triple=not conflict_candidate and record["fom_check"]["status"] == "conflict",
    )
    if conflict_candidate:
        metrics = {
            "reported_rate_gbps": record["speed"].value,
            "lane_rate_gbps": record["speed"].value,
            "power_mw": record["power"].value,
            "energy_pj_bit": record["energy"].value,
        }
    if not any(value is not None for value in metrics.values()):
        return None
    implementation_id = upsert_implementation(
        cur, f"ieee:{paper['article_number']}", paper,
        paper.get("title") or record["title"], int(paper.get("year") or record["year"]),
        paper.get("source_name") or record["publication"], "exact_title_year_venue", 1.0,
    )
    title_metrics = extract_performance(record["title"])
    lane_count = title_metrics.get("lane_count") or 1
    review_status = "needs_review" if conflict_candidate else "verified"
    confidence = 0.70 if conflict_candidate else 0.99
    measurement = {
        "implementation_id": implementation_id,
        "operating_point_key": operating_point_key(record, "-candidate" if conflict_candidate else ""),
        "reference_title": record["title"], "publication_year": record["year"],
        "publication_name": record["publication"], "first_author": record["first_author"],
        "affiliation": record["affiliation"], "link_class": link_class_for(record["title"]),
        "component_scope": component_scope_for(record["title"]),
        "modulation": modulation_for(record["title"]),
        "process_text": None if conflict_candidate else record["process_text"],
        "process_nm": metrics.get("process_nm"),
        "reported_rate_text": record["speed"].raw if metrics.get("reported_rate_gbps") is not None else None,
        "reported_rate_gbps": metrics.get("reported_rate_gbps"),
        "rate_scope": "lane" if metrics.get("lane_rate_gbps") is not None else "unknown",
        "lane_rate_gbps": metrics.get("lane_rate_gbps"),
        "lane_count": lane_count if metrics.get("lane_rate_gbps") is not None else None,
        "aggregate_rate_gbps": None, "aggregate_rate_basis": None,
        "power_mw": metrics.get("power_mw"), "power_scope": "unknown",
        "energy_pj_bit": metrics.get("energy_pj_bit"), "energy_scope": "reported",
        "energy_basis": "reported_user_sheet" if metrics.get("energy_pj_bit") is not None else None,
        "channel_loss_db": metrics.get("channel_loss_db"),
        "loss_frequency_ghz": metrics.get("loss_frequency_ghz"),
        "source_kind": "user_sheet", "overall_confidence": confidence,
        "review_status": review_status,
    }
    measurement_id = upsert_measurement(cur, measurement)
    field_columns = {
        "process_nm": "Process", "reported_rate_gbps": "Speed (Gb/s)",
        "lane_rate_gbps": "Speed (Gb/s)", "power_mw": "Power(mW)",
        "energy_pj_bit": "Energy/Bit [pJ]", "channel_loss_db": "Loss (dB)",
        "loss_frequency_ghz": "Loss Frequency (GHz)",
    }
    locator = f"'{record['sheet']}'!A{record['row']}:N{record['row']}"
    source_text = json.dumps(
        {"workbook_sha256": workbook_hash, "row": record["raw_record"]},
        ensure_ascii=False, sort_keys=True, default=str,
    )
    for field, column in field_columns.items():
        if measurement.get(field) is None:
            continue
        evidence_text = (
            f"{column}: {record['raw_record'].get(column)}; Title: {record['title']}; "
            f"workbook_sha256: {workbook_hash}"
        )
        upsert_evidence(
            cur, measurement_id, field, evidence_text, "user_sheet", record["source_url"],
            locator, paper["id"], None, "user_sheet_cross_validation", confidence,
            review_status, source_text, extractor_version=SYNC_VERSION,
        )
    return measurement_id


def audit_and_sync(conn, workbook_path: Path, *, threshold: float = 0.94,
                   margin: float = 0.025, apply: bool = False) -> dict:
    records = load_reference_rows(workbook_path)
    snapshot_hash = workbook_sha256(workbook_path)
    with conn.cursor() as cur:
        by_year_source, by_source_title = fetch_candidates(cur, records)
        screenings = fetch_screenings(cur)
        for record in records:
            record["match"] = match_record(record, by_year_source, by_source_title, threshold, margin)

        duplicate_groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
        for record in records:
            if record["sheet_duplicate"]:
                duplicate_groups[(record["sheet"], record["normalized_title"])].append(record)
        duplicate_ambiguous_ids = set()
        for group in duplicate_groups.values():
            matched_ids = [item["match"]["paper"]["id"] for item in group if item["match"].get("paper")]
            if len(matched_ids) != len(group) or len(set(matched_ids)) != len(group):
                duplicate_ambiguous_ids.update((item["sheet"], item["row"]) for item in group)

        paper_ids = sorted({record["match"]["paper"]["id"] for record in records if record["match"].get("paper")})
        existing_by_paper = fetch_existing_measurements(cur, paper_ids)
        report_rows = []
        counters = Counter()
        comparison_counts: dict[str, Counter] = defaultdict(Counter)

        for record in records:
            match = record["match"]
            paper = match.get("paper")
            duplicate_ambiguous = (record["sheet"], record["row"]) in duplicate_ambiguous_ids
            screening = screenings.get(int(paper["id"]), {}) if paper else {}
            high_confidence_match = paper is not None and match["method"] in {
                "exact_title_year_venue", "exact_title_venue_year_offset", "fuzzy_title_year_venue",
            } and not duplicate_ambiguous
            eligible = high_confidence_match and bool(screening.get("include_in_survey"))
            counters["records"] += 1
            counters[match["method"]] += 1
            counters["matched"] += int(paper is not None)
            counters["high_confidence_match"] += int(high_confidence_match)
            counters["eligible"] += int(eligible)
            counters["sheet_duplicate_ambiguous"] += int(duplicate_ambiguous)
            counters["internal_fom_conflict"] += int(record["fom_check"]["status"] == "conflict")
            fom_review_needed = (
                record["fom_check"]["status"] == "conflict"
                or any(field in record["uncertain_fields"] for field in (
                    "speed_gbps", "power_mw", "energy_pj_bit",
                ))
            )
            counters["fom_review_total"] += int(fom_review_needed)
            comparisons = {}
            if paper:
                counters["fixed_population"] += int(bool(screening.get("include_in_survey")))
                metrics = reference_metrics(record)
                existing_rows = existing_by_paper.get(int(paper["id"]), [])
                for field, value in metrics.items():
                    if value is None or field not in COMPARE_TOLERANCES:
                        continue
                    result = compare_field(value, best_existing_for_field(existing_rows, field), field)
                    comparisons[field] = result
                    comparison_counts[field][result["status"]] += 1
            applied_ids = []
            if apply and eligible:
                primary_id = store_record_measurement(cur, record, paper, snapshot_hash)
                if primary_id:
                    applied_ids.append(primary_id)
                if record["fom_check"]["status"] == "conflict":
                    candidate_id = store_record_measurement(
                        cur, record, paper, snapshot_hash, conflict_candidate=True,
                    )
                    if candidate_id:
                        applied_ids.append(candidate_id)
                counters["applied_records"] += 1
                counters["applied_measurements"] += len(applied_ids)

            report_rows.append({
                "sheet": record["sheet"], "gid": record["gid"], "row": record["row"],
                "year": record["year"], "publication": record["publication"],
                "title": record["title"], "main_category": record["main_category"],
                "match_method": match["method"], "match_score": match["score"],
                "match_margin": match["margin"],
                "article_number": paper.get("article_number") if paper else None,
                "paper_id": paper.get("id") if paper else None,
                "database_title": paper.get("title") if paper else None,
                "database_year": paper.get("year") if paper else None,
                "database_venue": paper.get("source_name") if paper else None,
                "relevance_class": screening.get("relevance_class"),
                "in_fixed_population": bool(screening.get("include_in_survey")),
                "eligible": eligible, "sheet_duplicate_ambiguous": duplicate_ambiguous,
                "process_nm": record["process_nm"], "process_qualifier": record["process_qualifier"],
                "speed_gbps": record["speed"].value, "power_mw": record["power"].value,
                "energy_pj_bit": record["energy"].value,
                "channel_loss_db": record["loss"].value,
                "loss_frequency_ghz": record["loss_frequency"].value,
                "fom_check": record["fom_check"], "comparisons": comparisons,
                "fom_review_needed": fom_review_needed,
                "uncertain_fields": record["uncertain_fields"],
                "applied_measurement_ids": applied_ids, "source_url": record["source_url"],
            })

        if apply:
            conn.commit()
        else:
            conn.rollback()

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_spreadsheet_id": SOURCE_SPREADSHEET_ID, "source_url": SOURCE_SPREADSHEET_URL,
        "snapshot_path": str(workbook_path.resolve()), "snapshot_sha256": snapshot_hash,
        "sync_version": SYNC_VERSION, "apply": apply, "threshold": threshold, "margin": margin,
        "summary": dict(counters),
        "records_by_sheet": dict(Counter(row["sheet"] for row in records)),
        "records_by_category": dict(Counter(row["main_category"] or "Unspecified" for row in records)),
        "comparison_summary": {field: dict(counts) for field, counts in comparison_counts.items()},
        "rows": report_rows,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--report", type=Path,
        default=ROOT / "outputs" / "serdes_reference" / SURVEY_REPORT_NAME,
    )
    parser.add_argument("--match-threshold", type=float, default=0.94)
    parser.add_argument("--match-margin", type=float, default=0.025)
    parser.add_argument("--apply", action="store_true")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    workbook = args.workbook if args.workbook.is_absolute() else ROOT / args.workbook
    report_path = args.report if args.report.is_absolute() else ROOT / args.report
    if not workbook.exists():
        raise SystemExit(f"Workbook not found: {workbook}")
    conn = db_connect()
    try:
        ensure_schema(conn)
        report = audit_and_sync(
            conn, workbook, threshold=args.match_threshold,
            margin=args.match_margin, apply=args.apply,
        )
        if args.apply:
            # A sheet resync can legitimately replace previously calculated
            # NULL fills. Rebuild the deterministic P/R layer immediately so
            # the SQL survey remains complete and provenance stays current.
            report["fom_regeneration_summary"] = regenerate_fom(
                conn, apply=True,
            )["summary"]
    finally:
        conn.close()
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "report": str(report_path), "snapshot_sha256": report["snapshot_sha256"],
        "apply": report["apply"], "summary": report["summary"],
        "comparison_summary": report["comparison_summary"],
        "fom_regeneration_summary": report.get("fom_regeneration_summary"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
