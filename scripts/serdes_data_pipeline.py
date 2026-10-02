# -*- coding: utf-8 -*-
r"""Build the structured SerDes survey layer without requiring a VPN.

The script is deliberately offline-first: Flask only reads cached SQL data.
IEEE abstracts are collected incrementally through the official Metadata API,
then immutable snapshots are parsed into traceable field-level evidence.

Examples:
  .venv\Scripts\python.exe scripts\serdes_data_pipeline.py --migrate-only
  .venv\Scripts\python.exe scripts\serdes_data_pipeline.py --import-wlink tmp\WLink_survey_2025.xlsx --extract-titles
  .venv\Scripts\python.exe scripts\serdes_data_pipeline.py --fetch-abstracts 40 --extract-abstracts
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import time

import openpyxl
import pymysql
import requests
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from serdes_metrics import (  # noqa: E402
    LINK_MEDIUM_VERSION,
    SERDES_SQL_PATTERN,
    SERDES_SCREENING_VERSION,
    canonical_ieee_url,
    classify_link_medium,
    extract_performance,
    normalize_title,
    plain_text,
    screen_serdes_relevance,
    taxonomy,
    title_similarity,
)
from serdes_enrichment import (  # noqa: E402
    ENERGY_COMPONENT_SCOPE_VERSION,
    LINK_SUBTYPE_VERSION,
    classify_energy_component_scope,
    classify_link_subtype,
)
from serdes_implementation_families import (  # noqa: E402
    build_implementation_families,
)

load_dotenv(ROOT / ".env")

MIGRATION_PATHS = (
    ROOT / "scripts" / "migrations" / "001_serdes_measurements.sql",
    ROOT / "scripts" / "migrations" / "002_serdes_screening.sql",
    ROOT / "scripts" / "migrations" / "003_metadata_fetch_state.sql",
    ROOT / "scripts" / "migrations" / "004_current_abstract.sql",
    ROOT / "scripts" / "migrations" / "005_serdes_link_medium.sql",
    ROOT / "scripts" / "migrations" / "006_serdes_enrichment.sql",
    ROOT / "scripts" / "migrations" / "007_serdes_implementation_families.sql",
    ROOT / "scripts" / "migrations" / "008_local_ai.sql",
    ROOT / "scripts" / "migrations" / "009_ai_recommendation_profiles.sql",
    ROOT / "scripts" / "migrations" / "010_serdes_bibliography.sql",
    ROOT / "scripts" / "migrations" / "011_recommendation_feedback.sql",
)
IEEE_API_URL = "https://ieeexploreapi.ieee.org/api/v1/search/articles"
WLINK_SOURCE_URL = (
    "https://web.engr.oregonstate.edu/~anandt/linksurvey/data/"
    "WLink_survey_2025.xlsx"
)
EXTRACTOR_VERSION = "serdes-regex-1.0"
SERDES_LATEST_COMPLETE_RUNS_SQL = (
    "SELECT MAX(latest_run.id) FROM serdes_screening_runs latest_run "
    "WHERE latest_run.status='complete' "
    # PyMySQL interpolates even an empty parameter tuple. Escape the literal
    # wildcard; unparameterized queries interpret the doubled wildcard equally.
    "AND latest_run.scope_name LIKE 'all%%' GROUP BY latest_run.venue"
)
MEASUREMENT_COLUMNS = (
    "implementation_id", "operating_point_key", "reference_title",
    "publication_year", "publication_name", "first_author", "affiliation",
    "link_class", "component_scope", "modulation", "process_text", "process_nm",
    "reported_rate_text", "reported_rate_gbps", "reported_rate_min_gbps",
    "reported_rate_max_gbps", "rate_scope", "lane_rate_gbps", "lane_count",
    "aggregate_rate_gbps", "aggregate_rate_basis", "symbol_rate_gbaud",
    "throughput_density_gbps_per_mm", "power_mw", "power_scope",
    "energy_pj_bit", "energy_scope", "energy_basis",
    "energy_loss_normalized_pj_bit_db", "channel_loss_db",
    "loss_frequency_ghz", "ber", "ber_scope", "active_area_mm2",
    "source_kind", "overall_confidence", "review_status",
)

PERFORMANCE_FIELDS = (
    "reported_rate_gbps", "reported_rate_min_gbps", "reported_rate_max_gbps",
    "lane_rate_gbps", "lane_count", "aggregate_rate_gbps", "symbol_rate_gbaud",
    "throughput_density_gbps_per_mm", "power_mw", "energy_pj_bit",
    "energy_loss_normalized_pj_bit_db", "process_nm", "channel_loss_db",
    "loss_frequency_ghz", "ber",
)


def db_connect():
    return pymysql.connect(
        host=os.getenv("DB_HOST", "127.0.0.1"),
        port=int(os.getenv("DB_PORT", "3306")),
        user=os.getenv("DB_USER", "root"),
        password=os.getenv("DB_PASSWORD", ""),
        database=os.getenv("DB_NAME", "ieee_repo"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=False,
    )


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def clean_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text.lower() in {"nan", "n/a", "na", "-"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def split_sql_statements(sql_text: str) -> list[str]:
    lines = []
    for line in sql_text.splitlines():
        if line.lstrip().startswith("--"):
            continue
        lines.append(line)
    return [statement.strip() for statement in "\n".join(lines).split(";") if statement.strip()]


def ensure_schema(conn) -> int:
    statements = []
    for path in MIGRATION_PATHS:
        statements.extend(split_sql_statements(path.read_text(encoding="utf-8")))
    with conn.cursor() as cur:
        for statement in statements:
            cur.execute(statement)
    conn.commit()
    return len(statements)


def component_scope_for(title_text: str) -> str:
    blocks = set(taxonomy(title_text)["blocks"])
    if "TX" in blocks and "RX" in blocks:
        return "tx_rx"
    if "TX" in blocks:
        return "tx"
    if "RX" in blocks:
        return "rx"
    return "full_link" if "System" in blocks else "unknown"


def modulation_for(title_text: str) -> str | None:
    signals = taxonomy(title_text)["signals"]
    return None if signals == ["Other"] else signals[0].lower().replace("-", "")


def link_class_for(title_text: str) -> str | None:
    value = plain_text(title_text).lower()
    rules = (
        ("die_to_die", r"die[- ]to[- ]die|chiplet|ucie"),
        ("chip_to_chip", r"chip[- ]to[- ]chip"),
        ("backplane", r"backplane"),
        ("cable", r"cable|copper"),
        ("memory", r"ddr|gddr|hbm|memory interface"),
        ("optical", r"optical|vcsel|silicon photonic|photodiode"),
    )
    for label, pattern in rules:
        if re.search(pattern, value):
            return label
    return None


def classify_link_media(conn, included_only: bool = False) -> dict:
    """Rebuild auditable paper-level Optical/Electrical/Unspecified labels."""
    scope_join = ""
    if included_only:
        scope_join = (
            "JOIN serdes_paper_screenings scope ON scope.paper_id=p.id "
            f"AND scope.run_id IN ({SERDES_LATEST_COMPLETE_RUNS_SQL}) "
            "AND scope.include_in_survey=1 "
        )
    with conn.cursor() as cur:
        cur.execute(
            "SELECT p.id, p.title, a.id abstract_id, a.abstract_text, "
            "screening.relevance_class "
            "FROM papers p "
            f"{scope_join}"
            "LEFT JOIN paper_current_abstracts a "
            "  ON a.paper_id=p.id AND a.is_current=1 "
            "LEFT JOIN serdes_paper_screenings screening ON screening.paper_id=p.id "
            f" AND screening.run_id IN ({SERDES_LATEST_COMPLETE_RUNS_SQL}) "
            "WHERE p.source_system='ieee'"
        )
        papers = cur.fetchall()

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    payloads = []
    stats = {"total": len(papers), "optical": 0, "electrical": 0, "unspecified": 0}
    source_stats = {}
    for paper in papers:
        decision = classify_link_medium(
            paper.get("title"),
            paper.get("abstract_text"),
            screening_class=paper.get("relevance_class"),
        )
        medium = decision["link_medium"]
        stats[medium] += 1
        source = decision["medium_source"]
        source_stats[source] = source_stats.get(source, 0) + 1
        payloads.append((
            int(paper["id"]), paper.get("abstract_id"), medium, source,
            decision["medium_confidence"],
            json.dumps(decision["medium_reason_codes"], ensure_ascii=False),
            decision.get("medium_evidence"), LINK_MEDIUM_VERSION, now,
        ))

    try:
        with conn.cursor() as cur:
            cur.executemany(
                """
                INSERT INTO serdes_paper_link_media
                    (paper_id, abstract_id, link_medium, medium_source,
                     medium_confidence, reason_codes, evidence_text,
                     classifier_version, classified_at)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                    abstract_id=VALUES(abstract_id),
                    link_medium=VALUES(link_medium),
                    medium_source=VALUES(medium_source),
                    medium_confidence=VALUES(medium_confidence),
                    reason_codes=VALUES(reason_codes),
                    evidence_text=VALUES(evidence_text),
                    classifier_version=VALUES(classifier_version),
                    classified_at=VALUES(classified_at)
                """,
                payloads,
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    stats["sources"] = source_stats
    stats["classifier_version"] = LINK_MEDIUM_VERSION
    return stats


def classify_measurement_scopes(conn) -> dict:
    """Rebuild operating-point energy circuit coverage without changing values."""
    # Earlier classifier versions wrote scope guesses for measurements without
    # an Energy/bit value. They are not comparable records; remove only those
    # generated rows while preserving human overrides in the separate table.
    with conn.cursor() as cur:
        cur.execute(
            "DELETE metric_scope FROM serdes_measurement_metric_scopes metric_scope "
            "JOIN serdes_measurements m ON m.id=metric_scope.measurement_id "
            "WHERE m.energy_pj_bit IS NULL"
        )
    conn.commit()
    with conn.cursor() as cur:
        cur.execute(
            "SELECT m.id measurement_id, m.component_scope, "
            "COALESCE(p.title, m.reference_title, i.canonical_title) title, "
            "(SELECT e.evidence_text FROM serdes_measurement_evidence e "
            " WHERE e.measurement_id=m.id "
            "   AND e.field_name IN ('energy_pj_bit','power_mw') "
            " ORDER BY "
            "   CASE e.review_status WHEN 'verified' THEN 0 WHEN 'reviewed' THEN 1 "
            "        WHEN 'curated_reference' THEN 2 ELSE 3 END, "
            "   CASE e.source_kind WHEN 'manual' THEN 0 WHEN 'pdf' THEN 1 "
            "        WHEN 'reference_xlsx' THEN 2 WHEN 'user_sheet' THEN 3 "
            "        WHEN 'abstract' THEN 4 WHEN 'title' THEN 5 ELSE 6 END, "
            "   CASE e.field_name WHEN 'energy_pj_bit' THEN 0 ELSE 1 END, e.id "
            " LIMIT 1) metric_evidence "
            "FROM serdes_measurements m "
            "JOIN serdes_implementations i ON i.id=m.implementation_id "
            "LEFT JOIN papers p ON p.id=i.canonical_paper_id "
            "WHERE m.energy_pj_bit IS NOT NULL"
        )
        rows = cur.fetchall()

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    stats = {"total": len(rows), **{value: 0 for value in (
        "tx", "rx", "trx", "full_link", "driver_only", "unknown",
    )}}
    payloads = []
    for row in rows:
        evidence_texts = [row["metric_evidence"]] if row.get("metric_evidence") else []
        decision = classify_energy_component_scope(
            row.get("title"), evidence_texts, row.get("component_scope"),
        )
        scope = decision["energy_component_scope"]
        stats[scope] += 1
        payloads.append((
            int(row["measurement_id"]), scope, decision["source"],
            decision["confidence"],
            json.dumps(decision["reason_codes"], ensure_ascii=False),
            decision.get("evidence"), ENERGY_COMPONENT_SCOPE_VERSION, now,
        ))
    try:
        with conn.cursor() as cur:
            cur.executemany(
                """
                INSERT INTO serdes_measurement_metric_scopes
                    (measurement_id, energy_component_scope, scope_source,
                     scope_confidence, reason_codes, evidence_text,
                     classifier_version, classified_at)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                    energy_component_scope=VALUES(energy_component_scope),
                    scope_source=VALUES(scope_source),
                    scope_confidence=VALUES(scope_confidence),
                    reason_codes=VALUES(reason_codes),
                    evidence_text=VALUES(evidence_text),
                    classifier_version=VALUES(classifier_version),
                    classified_at=VALUES(classified_at)
                """,
                payloads,
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    stats["classifier_version"] = ENERGY_COMPONENT_SCOPE_VERSION
    return stats


def classify_link_subtypes(conn, included_only: bool = False) -> dict:
    """Rebuild paper-level medium-specific subtypes with auditable evidence."""
    scope_join = ""
    if included_only:
        scope_join = (
            "JOIN serdes_paper_screenings scope ON scope.paper_id=p.id "
            f"AND scope.run_id IN ({SERDES_LATEST_COMPLETE_RUNS_SQL}) "
            "AND scope.include_in_survey=1 "
        )
    with conn.cursor() as cur:
        cur.execute(
            "SELECT p.id, p.title, a.abstract_text, medium.link_medium, "
            "medium.evidence_text medium_evidence, "
            "(SELECT m.link_class FROM serdes_implementation_papers ip "
            " JOIN serdes_measurements m ON m.implementation_id=ip.implementation_id "
            " WHERE ip.paper_id=p.id AND m.link_class IS NOT NULL "
            " ORDER BY FIELD(m.source_kind,'manual','user_sheet','pdf',"
            " 'reference_xlsx','abstract','title') LIMIT 1) stored_link_class "
            "FROM papers p "
            f"{scope_join}"
            "LEFT JOIN paper_current_abstracts a "
            "  ON a.paper_id=p.id AND a.is_current=1 "
            "LEFT JOIN serdes_paper_link_media medium ON medium.paper_id=p.id "
            "WHERE p.source_system='ieee'"
        )
        rows = cur.fetchall()

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    stats = {"total": len(rows)}
    payloads = []
    for row in rows:
        medium = str(row.get("link_medium") or "unspecified").lower()
        decision = classify_link_subtype(
            row.get("title"), row.get("abstract_text"), medium,
            row.get("stored_link_class"), row.get("medium_evidence"),
        )
        subtype = decision["link_subtype"]
        stats[subtype] = stats.get(subtype, 0) + 1
        payloads.append((
            int(row["id"]), medium, subtype, decision["source"],
            decision["confidence"],
            json.dumps(decision["reason_codes"], ensure_ascii=False),
            decision.get("evidence"), LINK_SUBTYPE_VERSION, now,
        ))
    try:
        with conn.cursor() as cur:
            cur.executemany(
                """
                INSERT INTO serdes_paper_link_subtypes
                    (paper_id, link_medium, link_subtype, subtype_source,
                     subtype_confidence, reason_codes, evidence_text,
                     classifier_version, classified_at)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                    link_medium=VALUES(link_medium),
                    link_subtype=VALUES(link_subtype),
                    subtype_source=VALUES(subtype_source),
                    subtype_confidence=VALUES(subtype_confidence),
                    reason_codes=VALUES(reason_codes),
                    evidence_text=VALUES(evidence_text),
                    classifier_version=VALUES(classifier_version),
                    classified_at=VALUES(classified_at)
                """,
                payloads,
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    stats["classifier_version"] = LINK_SUBTYPE_VERSION
    return stats


def process_values(value) -> tuple[str | None, float | None]:
    if value is None:
        return None, None
    text = plain_text(str(value))
    if not text:
        return None, None
    numbers = [float(item) for item in re.findall(r"\d+(?:\.\d+)?", text)]
    plausible = [item for item in numbers if 1 <= item <= 500]
    return text, min(plausible) if plausible else None


def upsert_implementation(
    cur,
    implementation_key: str,
    paper: dict | None,
    title: str,
    year: int | None,
    venue: str | None,
    dedup_method: str,
    confidence: float,
) -> int:
    paper_id = (paper.get("id") or paper.get("paper_id")) if paper else None
    cur.execute(
        """
        INSERT INTO serdes_implementations
            (implementation_key, canonical_paper_id, canonical_title,
             canonical_year, canonical_venue, dedup_method, dedup_confidence)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            id=LAST_INSERT_ID(id),
            canonical_paper_id=COALESCE(canonical_paper_id, VALUES(canonical_paper_id)),
            canonical_title=COALESCE(canonical_title, VALUES(canonical_title)),
            canonical_year=COALESCE(canonical_year, VALUES(canonical_year)),
            canonical_venue=COALESCE(canonical_venue, VALUES(canonical_venue)),
            dedup_method=IF(
                VALUES(dedup_method) IN
                    ('exact_title_year_venue','fuzzy_title_year_venue','unmatched_reference'),
                VALUES(dedup_method), dedup_method
            ),
            dedup_confidence=IF(
                VALUES(dedup_method) IN
                    ('exact_title_year_venue','fuzzy_title_year_venue','unmatched_reference'),
                VALUES(dedup_confidence), dedup_confidence
            )
        """,
        (implementation_key, paper_id, title, year, venue, dedup_method, confidence),
    )
    implementation_id = cur.lastrowid
    if paper_id:
        cur.execute(
            """
            INSERT INTO serdes_implementation_papers
                (implementation_id, paper_id, relation_type)
            VALUES (%s, %s, 'primary')
            ON DUPLICATE KEY UPDATE relation_type=VALUES(relation_type)
            """,
            (implementation_id, paper_id),
        )
    return implementation_id


def upsert_measurement(cur, values: dict) -> int:
    row = {column: values.get(column) for column in MEASUREMENT_COLUMNS}
    for field, default in (
        ("component_scope", "unknown"), ("rate_scope", "unknown"),
        ("power_scope", "unknown"), ("energy_scope", "unknown"),
        ("ber_scope", "unknown"), ("review_status", "extracted"),
    ):
        row[field] = row[field] or default
    row["overall_confidence"] = (
        0.7 if row["overall_confidence"] is None else row["overall_confidence"]
    )
    placeholders = ", ".join(["%s"] * len(MEASUREMENT_COLUMNS))
    assignments = ", ".join(
        f"{column}=VALUES({column})"
        for column in MEASUREMENT_COLUMNS
        if column not in {"implementation_id", "operating_point_key"}
    )
    cur.execute(
        f"""
        INSERT INTO serdes_measurements ({', '.join(MEASUREMENT_COLUMNS)})
        VALUES ({placeholders})
        ON DUPLICATE KEY UPDATE id=LAST_INSERT_ID(id), {assignments}
        """,
        tuple(row[column] for column in MEASUREMENT_COLUMNS),
    )
    return cur.lastrowid


def upsert_evidence(
    cur,
    measurement_id: int,
    field_name: str,
    evidence_text: str,
    source_kind: str,
    source_url: str | None,
    source_locator: str | None,
    paper_id: int | None,
    abstract_id: int | None,
    extraction_method: str,
    confidence: float,
    review_status: str,
    source_text: str,
    extractor_version: str = EXTRACTOR_VERSION,
):
    evidence_text = plain_text(evidence_text)
    evidence_sha = sha256_text(evidence_text)
    evidence_key = sha256_text(
        "|".join((str(measurement_id), field_name, source_kind, evidence_sha))
    )
    cur.execute(
        """
        INSERT INTO serdes_measurement_evidence
            (evidence_key, measurement_id, field_name, paper_id, abstract_id,
             source_kind, source_url, source_locator, evidence_text,
             source_sha256, evidence_sha256, extraction_method,
             extractor_version, confidence, review_status)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            source_url=VALUES(source_url), source_locator=VALUES(source_locator),
            confidence=VALUES(confidence), review_status=VALUES(review_status),
            abstract_id=VALUES(abstract_id), updated_at=CURRENT_TIMESTAMP
        """,
        (
            evidence_key, measurement_id, field_name, paper_id, abstract_id,
            source_kind, source_url, source_locator, evidence_text,
            sha256_text(plain_text(source_text)), evidence_sha, extraction_method,
            extractor_version, confidence, review_status,
        ),
    )


def performance_measurement(performance: dict, **base) -> dict:
    values = dict(base)
    for field in PERFORMANCE_FIELDS:
        values[field] = performance.get(field)
    values.update({
        "reported_rate_text": performance.get("reported_rate_text"),
        "rate_scope": performance.get("rate_scope") or "unknown",
        "aggregate_rate_basis": (
            "lane_x_count" if performance.get("aggregate_rate_gbps") is not None else None
        ),
        "energy_basis": None,
    })
    if performance.get("energy_pj_bit") is not None:
        energy_evidence = performance.get("evidence", {}).get("energy_pj_bit", {})
        values["energy_basis"] = (
            "converted_power_per_rate"
            if energy_evidence.get("converted_from")
            else "reported_energy"
        )
    if performance.get("ber_scope"):
        values["ber_scope"] = performance["ber_scope"]
    return values


def store_extracted_evidence(
    cur,
    measurement_id: int,
    performance: dict,
    source_kind: str,
    source_url: str,
    source_locator: str,
    paper_id: int,
    abstract_id: int | None,
    confidence: float,
    source_text: str,
):
    for field_name, item in performance.get("evidence", {}).items():
        upsert_evidence(
            cur, measurement_id, field_name, item.get("excerpt") or item.get("raw") or source_text,
            source_kind, source_url, source_locator, paper_id, abstract_id,
            "regex", confidence, "extracted", source_text,
        )

    candidate_specs = {
        "energy_pj_bit": ("energies", {"energy"}),
        "energy_loss_normalized_pj_bit_db": ("energies", {"loss_normalized_energy"}),
        "power_mw": ("powers", {"power_mw"}),
        "process_nm": ("processes", {"process_nm"}),
        "channel_loss_db": ("losses", {"channel_loss_db"}),
        "ber": ("ber", {"ber"}),
        "reported_rate_gbps": ("rates", {"bit_rate", "bit_rate_per_lane"}),
        "symbol_rate_gbaud": ("rates", {"symbol_rate"}),
    }
    ambiguous_fields = set(performance.get("ambiguous_fields", []))
    for field_name, (candidate_group, accepted_kinds) in candidate_specs.items():
        if field_name not in ambiguous_fields:
            continue
        for item in performance.get("candidates", {}).get(candidate_group, []):
            if item.get("kind") in accepted_kinds:
                upsert_evidence(
                    cur, measurement_id, f"{field_name}_candidate",
                    item.get("excerpt") or item["raw"], source_kind, source_url,
                    source_locator, paper_id, abstract_id, "regex_candidate",
                    max(0.4, confidence - 0.2), "needs_review", source_text,
                )


def fetch_wlink_candidates(cur) -> dict[tuple[int, str], list[dict]]:
    cur.execute(
        """
        SELECT id, article_number, title, authors, year, source_name
        FROM papers
        WHERE source_system='ieee'
          AND CAST(year AS UNSIGNED) BETWEEN 2000 AND 2026
          AND source_name IN
              ('ISSCC','VLSI-Circuits','VLSI-Tech','CICC','ASSCC','ESSCIRC','JSSC')
        """
    )
    by_year_source: dict[tuple[int, str], list[dict]] = {}
    for row in cur.fetchall():
        try:
            year = int(row.get("year") or 0)
        except (TypeError, ValueError):
            continue
        source = str(row.get("source_name") or "").upper()
        by_year_source.setdefault((year, source), []).append(row)
    return by_year_source


def venue_aliases(value: str) -> tuple[str, ...]:
    venue = plain_text(value).strip().upper()
    aliases = {
        "VLSI": ("VLSI-CIRCUITS", "VLSI-TECH"),
        "ISSCC": ("ISSCC", "JSSC"),
        "A-SSCC": ("ASSCC",),
        "ASSCC": ("ASSCC",),
    }
    return aliases.get(venue, (venue,))


def match_wlink_paper(
    title: str,
    year: int,
    venue: str,
    candidates: dict[tuple[int, str], list[dict]],
    threshold: float,
) -> tuple[dict | None, str, float]:
    pool = []
    for alias in venue_aliases(venue):
        pool.extend(candidates.get((year, alias), []))
    expected = normalize_title(title)
    exact = [row for row in pool if normalize_title(row.get("title")) == expected]
    if len(exact) == 1:
        return exact[0], "exact_title_year_venue", 1.0

    best = None
    best_score = 0.0
    for row in pool:
        score = title_similarity(title, row.get("title"))
        if score > best_score:
            best, best_score = row, score
    if best is not None and best_score >= threshold:
        return best, "fuzzy_title_year_venue", best_score
    return None, "unmatched_reference", best_score


def import_wlink(conn, workbook_path: Path, threshold: float = 0.86) -> dict:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["SurveyData"]
    rows = sheet.iter_rows(values_only=True)
    headers = [plain_text(value) for value in next(rows)]
    stats = {"read": 0, "exact": 0, "fuzzy": 0, "unmatched": 0, "measurements": 0}
    with conn.cursor() as cur:
        candidates = fetch_wlink_candidates(cur)
        for row_number, values in enumerate(rows, start=2):
            record = dict(zip(headers, values))
            title = plain_text(record.get("Title"))
            if not title:
                continue
            stats["read"] += 1
            year_number = clean_number(record.get("Year"))
            year = int(year_number) if year_number else None
            venue = plain_text(record.get("Publication"))
            paper, method, score = match_wlink_paper(
                title, year or 0, venue, candidates, threshold,
            )
            if method.startswith("exact"):
                stats["exact"] += 1
            elif method.startswith("fuzzy"):
                stats["fuzzy"] += 1
            else:
                stats["unmatched"] += 1

            if paper:
                implementation_key = f"ieee:{paper['article_number']}"
            else:
                implementation_key = f"wlink:{year}:{sha256_text(normalize_title(title))[:20]}"
            implementation_id = upsert_implementation(
                cur, implementation_key, paper, title, year, venue, method,
                0.95 if paper and score >= 0.95 else max(0.6, score),
            )

            title_metrics = extract_performance(title)
            speed = clean_number(record.get("Speed (Gb/s)"))
            lane_count = title_metrics.get("lane_count") or 1
            aggregate = speed * lane_count if speed else None
            process_text, process_nm = process_values(record.get("Process"))
            measurement = {
                "implementation_id": implementation_id,
                "operating_point_key": f"wlink-2025-row-{row_number}",
                "reference_title": title,
                "publication_year": year,
                "publication_name": venue,
                "first_author": plain_text(record.get("First Author")),
                "affiliation": plain_text(record.get("Affiliation")),
                "link_class": link_class_for(title),
                "component_scope": component_scope_for(title),
                "modulation": modulation_for(title),
                "process_text": process_text,
                "process_nm": process_nm,
                "reported_rate_text": str(record.get("Speed (Gb/s)")) if speed else None,
                "reported_rate_gbps": speed,
                "rate_scope": "lane",
                "lane_rate_gbps": speed,
                "lane_count": lane_count,
                "aggregate_rate_gbps": aggregate,
                "aggregate_rate_basis": "lane_x_count" if lane_count > 1 else "reported",
                "power_mw": clean_number(record.get("Power(mW)")),
                "power_scope": "unknown",
                "energy_pj_bit": clean_number(record.get("Energy/Bit [pJ]")),
                "energy_scope": "reported",
                "energy_basis": "reported",
                "channel_loss_db": clean_number(record.get("Loss (dB)")),
                "loss_frequency_ghz": clean_number(record.get("Loss Frequency (GHz)")),
                "source_kind": "reference_xlsx",
                "overall_confidence": 0.95,
                "review_status": "curated_reference",
            }
            measurement_id = upsert_measurement(cur, measurement)
            paper_id = paper.get("id") if paper else None
            locator = f"SurveyData!A{row_number}:L{row_number}"
            field_map = {
                "process_nm": "Process", "channel_loss_db": "Loss (dB)",
                "loss_frequency_ghz": "Loss Frequency (GHz)",
                "lane_rate_gbps": "Speed (Gb/s)", "aggregate_rate_gbps": "Speed (Gb/s)",
                "power_mw": "Power(mW)", "energy_pj_bit": "Energy/Bit [pJ]",
            }
            row_source = json.dumps(record, ensure_ascii=False, default=str)
            for field_name, column_name in field_map.items():
                if measurement.get(field_name) is None:
                    continue
                evidence = f"{column_name}: {record.get(column_name)}; Title: {title}"
                upsert_evidence(
                    cur, measurement_id, field_name, evidence, "reference_xlsx",
                    WLINK_SOURCE_URL, locator, paper_id, None, "survey_import",
                    0.95, "curated_reference", row_source,
                )
            stats["measurements"] += 1
            if stats["measurements"] % 50 == 0:
                conn.commit()
        conn.commit()
    workbook.close()
    return stats


def fetch_serdes_papers(cur, screened_only: bool = False) -> list[dict]:
    screening_join = (
        "JOIN serdes_paper_screenings s ON s.paper_id=p.id "
        f"AND s.run_id IN ({SERDES_LATEST_COMPLETE_RUNS_SQL}) "
        "AND s.include_in_survey=1"
        if screened_only else ""
    )
    scope_clause = (
        "" if screened_only
        else "AND LOWER(CONCAT_WS(' ', p.title, p.source_name, p.issue)) REGEXP %s"
    )
    params = () if screened_only else (SERDES_SQL_PATTERN,)
    cur.execute(
        f"""
        SELECT id, article_number, title, authors, year, source_name, url
        FROM papers p
        {screening_join}
        WHERE p.source_system='ieee'
          {scope_clause}
        """,
        params,
    )
    return cur.fetchall()


def extract_titles(conn, screened_only: bool = False) -> dict:
    stats = {"papers": 0, "measurements": 0, "evidence": 0}
    with conn.cursor() as cur:
        papers = fetch_serdes_papers(cur, screened_only=screened_only)
        for paper in papers:
            stats["papers"] += 1
            performance = extract_performance(paper.get("title"))
            has_metric = any(performance.get(field) is not None for field in PERFORMANCE_FIELDS)
            has_candidates = any(performance["candidates"].values())
            if not has_metric and not has_candidates:
                cur.execute(
                    "DELETE m FROM serdes_measurements m "
                    "JOIN serdes_implementations i ON i.id=m.implementation_id "
                    "WHERE i.implementation_key=%s "
                    "AND m.operating_point_key='title-clue' AND m.source_kind='title'",
                    (f"ieee:{paper['article_number']}",),
                )
                continue
            try:
                year = int(paper.get("year") or 0) or None
            except (TypeError, ValueError):
                year = None
            implementation_id = upsert_implementation(
                cur, f"ieee:{paper['article_number']}", paper, paper["title"],
                year, paper.get("source_name"), "single_paper", 1.0,
            )
            values = performance_measurement(
                performance,
                implementation_id=implementation_id,
                operating_point_key="title-clue",
                reference_title=paper.get("title"),
                publication_year=year,
                publication_name=paper.get("source_name"),
                link_class=link_class_for(paper.get("title")),
                component_scope=component_scope_for(paper.get("title")),
                modulation=modulation_for(paper.get("title")),
                power_scope="unknown", energy_scope="unknown", ber_scope="unknown",
                source_kind="title", overall_confidence=0.62,
                review_status="extracted",
            )
            measurement_id = upsert_measurement(cur, values)
            cur.execute(
                "DELETE FROM serdes_measurement_evidence "
                "WHERE measurement_id=%s AND source_kind='title'",
                (measurement_id,),
            )
            before = cur.rowcount
            store_extracted_evidence(
                cur, measurement_id, performance, "title",
                canonical_ieee_url(paper["article_number"], paper.get("url")),
                "Title", paper["id"], None, 0.62, paper["title"],
            )
            stats["measurements"] += 1
            stats["evidence"] += max(0, cur.rowcount, before)
            if stats["measurements"] % 250 == 0:
                conn.commit()
        conn.commit()
    return stats


class IEEEAbstractClient:
    def __init__(self, api_key: str, delay: float = 1.1):
        if not api_key:
            raise RuntimeError("IEEE_API_KEY is not configured in .env")
        self.api_key = api_key
        self.delay = max(0.0, delay)
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": "IEEE-Paper-Server-SerDes/1.0"})

    def fetch_batch(self, article_numbers: list[str]) -> dict[str, dict]:
        expected = {str(item) for item in article_numbers}
        params = {
            "apikey": self.api_key,
            "format": "json",
            "max_records": min(200, max(25, len(expected) * 4)),
            "querytext": " OR ".join(sorted(expected)),
        }
        for attempt in range(1, 5):
            try:
                response = self.session.get(IEEE_API_URL, params=params, timeout=90)
            except requests.RequestException as exc:
                if attempt == 4:
                    detail = str(exc).replace(self.api_key, "<redacted>")
                    raise RuntimeError(
                        f"IEEE Metadata API request failed: {detail}"
                    ) from None
                time.sleep(2**attempt)
                continue
            if response.status_code in {401, 403}:
                detail = plain_text(response.text)[:240].replace(self.api_key, "<redacted>")
                detail_lower = detail.lower()
                if response.status_code == 403 and any(
                    marker in detail_lower for marker in ("rate", "quota", "limit")
                ):
                    raise RuntimeError(
                        f"IEEE Metadata API rate/quota limit reached (403): {detail or 'no detail'}"
                    )
                raise RuntimeError(
                    f"IEEE Metadata API authorization failed ({response.status_code}): "
                    f"{detail or 'no detail'}"
                )
            if response.status_code == 429:
                retry_after = response.headers.get("Retry-After")
                try:
                    delay = max(1.0, float(retry_after))
                except (TypeError, ValueError):
                    delay = float(2**attempt)
                if attempt == 4:
                    raise RuntimeError("IEEE Metadata API quota/rate limit reached")
                time.sleep(min(60.0, delay))
                continue
            if response.status_code >= 500:
                if attempt == 4:
                    raise RuntimeError(f"IEEE Metadata API server error ({response.status_code})")
                time.sleep(2**attempt)
                continue
            if response.status_code != 200:
                raise RuntimeError(f"IEEE Metadata API HTTP {response.status_code}")
            payload = response.json()
            results = {}
            for article in payload.get("articles") or []:
                key = str(article.get("article_number") or "")
                if key in expected:
                    results[key] = article
            if self.delay:
                time.sleep(self.delay)
            return results
        return {}


def store_abstract_snapshot(
    cur,
    paper: dict,
    article: dict,
    provider: str = "ieee_metadata",
) -> tuple[int, bool]:
    abstract = plain_text(article.get("abstract"))
    content_hash = sha256_text(abstract)
    provider = plain_text(provider) or "ieee_metadata"
    cur.execute(
        """
        SELECT id, content_sha256 FROM paper_abstracts
        WHERE paper_id=%s AND provider=%s AND is_current=1
        """,
        (paper["id"], provider),
    )
    current = cur.fetchone()
    if current and current["content_sha256"] == content_hash:
        cur.execute(
            "UPDATE paper_abstracts SET retrieved_at=%s WHERE id=%s",
            (datetime.now(timezone.utc).replace(tzinfo=None), current["id"]),
        )
        return current["id"], False

    if current:
        cur.execute("UPDATE paper_abstracts SET is_current=0 WHERE id=%s", (current["id"],))
    source_url = plain_text(article.get("abstract_url")) or canonical_ieee_url(paper["article_number"])
    provider_record_id = plain_text(article.get("provider_record_id")) or paper["article_number"]
    cur.execute(
        """
        INSERT INTO paper_abstracts
            (paper_id, provider, provider_record_id, source_url, abstract_text,
             content_sha256, retrieved_at, is_current)
        VALUES (%s, %s, %s, %s, %s, %s, %s, 1)
        ON DUPLICATE KEY UPDATE
            id=LAST_INSERT_ID(id), is_current=1, retrieved_at=VALUES(retrieved_at),
            source_url=VALUES(source_url)
        """,
        (
            paper["id"], provider, provider_record_id, source_url, abstract,
            content_hash, datetime.now(timezone.utc).replace(tzinfo=None),
        ),
    )
    return cur.lastrowid, True


def upsert_fetch_state(
    cur,
    paper_id: int,
    status: str,
    detail: str | None = None,
):
    """Record one exact-ID metadata attempt without hiding retryable errors."""
    attempted_at = datetime.now(timezone.utc).replace(tzinfo=None)
    completed_at = attempted_at if status in {"abstract", "no_abstract", "not_found"} else None
    cur.execute(
        """
        INSERT INTO paper_metadata_fetch_state
            (paper_id, provider, fetch_status, attempt_count, detail,
             last_attempt_at, completed_at)
        VALUES (%s, 'ieee_metadata', %s, 1, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            fetch_status=VALUES(fetch_status),
            attempt_count=attempt_count+1,
            detail=VALUES(detail),
            last_attempt_at=VALUES(last_attempt_at),
            completed_at=VALUES(completed_at)
        """,
        (paper_id, status, (detail or "")[:500] or None, attempted_at, completed_at),
    )


def fetch_abstracts(
    conn,
    limit: int,
    batch_size: int,
    delay: float,
    venue: str | None = None,
    scope: str = "candidate",
    year_from: int | None = None,
    year_to: int | None = None,
    retry_terminal: bool = False,
) -> dict:
    if scope not in {"candidate", "venue", "screened"}:
        raise ValueError(f"Unsupported abstract scope: {scope}")
    client = IEEEAbstractClient((os.getenv("IEEE_API_KEY") or "").strip(), delay)
    venue_clause = "AND p.source_name=%s" if venue else ""
    year_from_clause = "AND CAST(p.year AS UNSIGNED)>=%s" if year_from else ""
    year_to_clause = "AND CAST(p.year AS UNSIGNED)<=%s" if year_to else ""
    terminal_clause = "" if retry_terminal else (
        "AND (fs.fetch_status IS NULL OR "
        "fs.fetch_status NOT IN ('no_abstract','not_found'))"
    )
    screening_join = (
        "JOIN serdes_paper_screenings s ON s.paper_id=p.id "
        f"AND s.run_id IN ({SERDES_LATEST_COMPLETE_RUNS_SQL})"
        if scope == "screened" else ""
    )
    scope_clause = ""
    if scope == "candidate":
        scope_clause = (
            "AND LOWER(CONCAT_WS(' ', p.title, p.source_name, p.issue)) REGEXP %s"
        )
    elif scope == "screened":
        scope_clause = (
            "AND s.relevance_class IN ('core','adjacent','needs_review')"
        )
    query_params = []
    if scope == "candidate":
        query_params.append(SERDES_SQL_PATTERN)
    if venue:
        query_params.append(venue)
    if year_from:
        query_params.append(year_from)
    if year_to:
        query_params.append(year_to)
    query_params.append(max(1, int(limit)))
    with conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT p.id, p.article_number, p.title, p.year, p.source_name, p.url,
                   MAX(CASE WHEN a.is_current=1 THEN 1 ELSE 0 END) AS has_abstract
            FROM papers p
            {screening_join}
            LEFT JOIN paper_abstracts a
              ON a.paper_id=p.id AND a.provider='ieee_metadata'
            LEFT JOIN paper_metadata_fetch_state fs
              ON fs.paper_id=p.id AND fs.provider='ieee_metadata'
            WHERE p.source_system='ieee'
              AND p.article_number REGEXP '^[0-9]+$'
              {scope_clause}
              {venue_clause}
              {year_from_clause}
              {year_to_clause}
              {terminal_clause}
            GROUP BY p.id, p.article_number, p.title, p.year, p.source_name, p.url
            HAVING has_abstract=0
            ORDER BY p.is_favorite DESC, CAST(p.year AS UNSIGNED) DESC, p.id DESC
            LIMIT %s
            """,
            tuple(query_params),
        )
        targets = cur.fetchall()

    stats = {
        "venue": venue or "all",
        "scope": scope,
        "year_from": year_from,
        "year_to": year_to,
        "targeted": len(targets),
        "fetched": 0,
        "new_revisions": 0,
        "no_abstract": 0,
        "not_found": 0,
        "fallback_requests": 0,
    }

    def persist_article(cur, paper, article):
        if not article:
            stats["not_found"] += 1
            upsert_fetch_state(cur, paper["id"], "not_found")
            return
        if not plain_text(article.get("abstract")):
            stats["no_abstract"] += 1
            upsert_fetch_state(cur, paper["id"], "no_abstract")
            return
        _, created = store_abstract_snapshot(cur, paper, article)
        stats["fetched"] += 1
        stats["new_revisions"] += int(created)
        upsert_fetch_state(cur, paper["id"], "abstract")

    for start in range(0, len(targets), max(1, batch_size)):
        batch = targets[start:start + max(1, batch_size)]
        article_numbers = [str(row["article_number"]) for row in batch]
        try:
            records = client.fetch_batch(article_numbers)
        except Exception as exc:
            detail = str(exc).replace(client.api_key, "<redacted>")
            with conn.cursor() as cur:
                for paper in batch:
                    upsert_fetch_state(cur, paper["id"], "error", detail)
            conn.commit()
            raise

        # Persist every primary-query hit before exact-ID fallbacks. If a later
        # fallback is rate-limited, completed records remain checkpointed and
        # only the exact ID that failed is marked retryable.
        with conn.cursor() as cur:
            for paper in batch:
                article = records.get(str(paper["article_number"]))
                if article:
                    persist_article(cur, paper, article)
        conn.commit()

        missing = [
            paper for paper in batch
            if str(paper["article_number"]) not in records
        ]
        for paper in missing:
            stats["fallback_requests"] += 1
            number = str(paper["article_number"])
            try:
                exact_records = client.fetch_batch([number])
            except Exception as exc:
                detail = str(exc).replace(client.api_key, "<redacted>")
                with conn.cursor() as cur:
                    upsert_fetch_state(cur, paper["id"], "error", detail)
                conn.commit()
                raise
            with conn.cursor() as cur:
                persist_article(cur, paper, exact_records.get(number))
            conn.commit()
        print(
            f"  abstracts {min(start + len(batch), len(targets)):,}/{len(targets):,} "
            f"(saved={stats['fetched']:,}, no_abstract={stats['no_abstract']:,}, "
            f"not_found={stats['not_found']:,})"
        )
    return stats


def extract_abstracts(
    conn,
    venue: str | None = None,
    screened_only: bool = False,
    year_from: int | None = None,
    year_to: int | None = None,
) -> dict:
    stats = {"abstracts": 0, "measurements": 0, "with_energy": 0, "ambiguous_energy": 0}
    screening_join = (
        "JOIN serdes_paper_screenings s ON s.paper_id=p.id "
        f"AND s.run_id IN ({SERDES_LATEST_COMPLETE_RUNS_SQL}) "
        "AND s.include_in_survey=1"
        if screened_only else ""
    )
    venue_clause = "AND p.source_name=%s" if venue else ""
    year_from_clause = "AND CAST(p.year AS UNSIGNED)>=%s" if year_from else ""
    year_to_clause = "AND CAST(p.year AS UNSIGNED)<=%s" if year_to else ""
    params = []
    if venue:
        params.append(venue)
    if year_from:
        params.append(year_from)
    if year_to:
        params.append(year_to)
    with conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT a.id AS abstract_id, a.paper_id, a.abstract_text, a.source_url,
                   p.article_number, p.title, p.year, p.source_name, p.url
            FROM paper_current_abstracts a
            JOIN papers p ON p.id=a.paper_id
            {screening_join}
            WHERE a.is_current=1 AND p.source_system='ieee'
              {venue_clause}
              {year_from_clause}
              {year_to_clause}
            ORDER BY a.id
            """,
            tuple(params),
        )
        rows = cur.fetchall()
        for row in rows:
            stats["abstracts"] += 1
            abstract_text = row["abstract_text"]
            abstract_performance = extract_performance(abstract_text)
            title_performance = extract_performance(row["title"])
            performance = dict(abstract_performance)
            performance["evidence"] = dict(abstract_performance.get("evidence", {}))
            title_fallback_fields = []
            rate_fields = (
                "reported_rate_text", "reported_rate_gbps", "reported_rate_min_gbps",
                "reported_rate_max_gbps", "rate_scope", "lane_rate_gbps", "lane_count",
                "aggregate_rate_gbps", "symbol_rate_gbaud",
                "throughput_density_gbps_per_mm", "process_nm",
            )
            for field in rate_fields:
                if (
                    performance.get(field) is None
                    and field not in performance.get("ambiguous_fields", [])
                    and title_performance.get(field) is not None
                ):
                    performance[field] = title_performance[field]
                    title_fallback_fields.append(field)
            try:
                year = int(row.get("year") or 0) or None
            except (TypeError, ValueError):
                year = None
            implementation_id = upsert_implementation(
                cur, f"ieee:{row['article_number']}", row, row["title"], year,
                row.get("source_name"), "single_paper", 1.0,
            )
            values = performance_measurement(
                performance,
                implementation_id=implementation_id,
                operating_point_key="abstract-current",
                reference_title=row.get("title"), publication_year=year,
                publication_name=row.get("source_name"),
                link_class=link_class_for(f"{row['title']}. {abstract_text}"),
                component_scope=component_scope_for(f"{row['title']}. {abstract_text}"),
                modulation=modulation_for(f"{row['title']}. {abstract_text}"),
                power_scope="unknown", energy_scope="unknown", ber_scope="unknown",
                source_kind="abstract", overall_confidence=0.80,
                review_status="extracted",
            )
            measurement_id = upsert_measurement(cur, values)
            cur.execute(
                "DELETE FROM serdes_measurement_evidence "
                "WHERE measurement_id=%s AND source_kind IN ('abstract','title')",
                (measurement_id,),
            )
            store_extracted_evidence(
                cur, measurement_id, abstract_performance, "abstract", row["source_url"],
                "Abstract", row["paper_id"], row["abstract_id"], 0.80, abstract_text,
            )
            for field in title_fallback_fields:
                evidence_key = (
                    "reported_rate_gbps" if field in {
                        "reported_rate_text", "lane_rate_gbps", "lane_count",
                        "aggregate_rate_gbps", "rate_scope",
                    } else field
                )
                item = title_performance.get("evidence", {}).get(evidence_key)
                if not item:
                    continue
                upsert_evidence(
                    cur, measurement_id, field, item.get("excerpt") or item.get("raw"),
                    "title", canonical_ieee_url(row["article_number"], row.get("url")),
                    "Title", row["paper_id"], None, "regex_fallback", 0.62,
                    "extracted", row["title"],
                )
            stats["measurements"] += 1
            stats["with_energy"] += int(abstract_performance.get("energy_pj_bit") is not None)
            stats["ambiguous_energy"] += int(
                "energy_pj_bit" in abstract_performance["ambiguous_fields"]
            )
            if stats["measurements"] % 100 == 0:
                conn.commit()
        conn.commit()
    return stats


def screen_venue(
    conn,
    venue: str,
    scope: str = "candidate",
    year_from: int | None = None,
    year_to: int | None = None,
) -> dict:
    """Persist one reproducible screening decision for every paper in scope."""
    if scope not in {"candidate", "all"}:
        raise ValueError(f"Unsupported screening scope: {scope}")
    scope_clause = (
        "AND LOWER(CONCAT_WS(' ', p.title, p.source_name, p.issue)) REGEXP %s"
        if scope == "candidate" else ""
    )
    params = [venue]
    if scope == "candidate":
        params.append(SERDES_SQL_PATTERN)
    year_from_clause = "AND CAST(p.year AS UNSIGNED)>=%s" if year_from else ""
    year_to_clause = "AND CAST(p.year AS UNSIGNED)<=%s" if year_to else ""
    if year_from:
        params.append(year_from)
    if year_to:
        params.append(year_to)
    scope_name = scope
    if year_from or year_to:
        scope_name = f"{scope}:{year_from or '*'}-{year_to or '*'}"
    with conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT p.id, p.article_number, p.title, p.year, p.source_name,
                   a.id AS abstract_id, a.abstract_text
            FROM papers p
            LEFT JOIN paper_current_abstracts a
              ON a.paper_id=p.id
            WHERE p.source_system='ieee' AND p.source_name=%s
              {scope_clause}
              {year_from_clause}
              {year_to_clause}
            ORDER BY p.id
            """,
            tuple(params),
        )
        rows = cur.fetchall()
        started_at = datetime.now(timezone.utc).replace(tzinfo=None)
        cur.execute(
            """
            INSERT INTO serdes_screening_runs
                (venue, scope_name, scope_version, target_count, status, started_at)
            VALUES (%s, %s, %s, %s, 'running', %s)
            """,
            (venue, scope_name, SERDES_SCREENING_VERSION, len(rows), started_at),
        )
        run_id = cur.lastrowid
    conn.commit()

    stats = {
        "run_id": run_id,
        "venue": venue,
        "scope": scope,
        "year_from": year_from,
        "year_to": year_to,
        "scope_version": SERDES_SCREENING_VERSION,
        "targeted": len(rows),
        "abstracts": 0,
        "core": 0,
        "adjacent": 0,
        "needs_review": 0,
        "out_of_scope": 0,
    }
    try:
        with conn.cursor() as cur:
            for index, row in enumerate(rows, 1):
                decision = screen_serdes_relevance(
                    row.get("title"), row.get("abstract_text")
                )
                relevance_class = decision["relevance_class"]
                stats[relevance_class] += 1
                stats["abstracts"] += int(bool(row.get("abstract_id")))
                review_status = (
                    "needs_review" if relevance_class == "needs_review"
                    else "auto_screened"
                )
                cur.execute(
                    """
                    INSERT INTO serdes_paper_screenings
                        (paper_id, run_id, abstract_id, venue, scope_version,
                         relevance_class, relevance_score, include_in_survey,
                         reason_codes, rationale, screening_source, review_status,
                         screened_at)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                        run_id=VALUES(run_id), abstract_id=VALUES(abstract_id),
                        venue=VALUES(venue), scope_version=VALUES(scope_version),
                        relevance_class=VALUES(relevance_class),
                        relevance_score=VALUES(relevance_score),
                        include_in_survey=VALUES(include_in_survey),
                        reason_codes=VALUES(reason_codes), rationale=VALUES(rationale),
                        screening_source=VALUES(screening_source),
                        review_status=VALUES(review_status), screened_at=VALUES(screened_at)
                    """,
                    (
                        row["id"], run_id, row.get("abstract_id"), venue,
                        decision["scope_version"], relevance_class,
                        decision["relevance_score"],
                        int(decision["include_in_survey"]),
                        json.dumps(decision["reason_codes"], ensure_ascii=False),
                        decision["rationale"][:1000],
                        decision["screening_source"], review_status,
                        datetime.now(timezone.utc).replace(tzinfo=None),
                    ),
                )
            completed_at = datetime.now(timezone.utc).replace(tzinfo=None)
            cur.execute(
                """
                UPDATE serdes_screening_runs
                SET core_count=%s, adjacent_count=%s, review_count=%s,
                    excluded_count=%s, abstract_count=%s,
                    status='complete', completed_at=%s
                WHERE id=%s
                """,
                (
                    stats["core"], stats["adjacent"], stats["needs_review"],
                    stats["out_of_scope"], stats["abstracts"], completed_at, run_id,
                ),
            )
        conn.commit()
    except Exception:
        conn.rollback()
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE serdes_screening_runs SET status='failed' WHERE id=%s",
                (run_id,),
            )
        conn.commit()
        raise
    return stats


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--migrate-only", action="store_true")
    parser.add_argument(
        "--classify-media", action="store_true",
        help="Rebuild auditable Optical/Electrical/Unspecified paper labels",
    )
    parser.add_argument(
        "--media-included-only", action="store_true",
        help="With --classify-media, limit the rebuild to Core + Adjacent papers",
    )
    parser.add_argument(
        "--classify-taxonomy", action="store_true",
        help="Rebuild energy component scopes and medium-specific link subtypes",
    )
    parser.add_argument(
        "--build-implementation-families", action="store_true",
        help=(
            "Build non-destructive conference/journal candidates and strict "
            "reciprocal-best implementation families"
        ),
    )
    parser.add_argument(
        "--family-dry-run", action="store_true",
        help="Classify implementation-family candidates without writing them",
    )
    parser.add_argument("--import-wlink", type=Path)
    parser.add_argument("--match-threshold", type=float, default=0.86)
    parser.add_argument("--extract-titles", action="store_true")
    parser.add_argument("--fetch-abstracts", type=int, metavar="LIMIT")
    parser.add_argument("--extract-abstracts", action="store_true")
    parser.add_argument("--extract-venue", metavar="VENUE")
    parser.add_argument("--screened-only", action="store_true")
    parser.add_argument("--venue", help="Limit abstract acquisition to one exact venue")
    parser.add_argument(
        "--abstract-scope", choices=("candidate", "screened", "venue"), default="candidate",
        help=(
            "Fetch title-regex candidates, the current screened shortlist, "
            "or every paper in --venue"
        ),
    )
    parser.add_argument("--screen-venue", metavar="VENUE")
    parser.add_argument(
        "--screen-scope", choices=("candidate", "all"), default="candidate",
    )
    parser.add_argument("--batch-size", type=int, default=20)
    parser.add_argument("--delay", type=float, default=1.1)
    parser.add_argument("--year-from", type=int)
    parser.add_argument("--year-to", type=int)
    parser.add_argument(
        "--retry-terminal",
        action="store_true",
        help="Retry exact IDs previously confirmed as no-abstract/not-found",
    )
    return parser


def main():
    parser = build_parser()
    args = parser.parse_args()
    for name, value in (("--year-from", args.year_from), ("--year-to", args.year_to)):
        if value is not None and not 1900 <= value <= datetime.now().year + 2:
            parser.error(f"{name} must be between 1900 and {datetime.now().year + 2}")
    if args.year_from and args.year_to and args.year_from > args.year_to:
        parser.error("--year-from must be less than or equal to --year-to")
    if args.abstract_scope in {"venue", "screened"} and args.fetch_abstracts is not None and not args.venue:
        parser.error(f"--abstract-scope {args.abstract_scope} requires --venue")
    if args.family_dry_run and not args.build_implementation_families:
        parser.error("--family-dry-run requires --build-implementation-families")
    actions = any((
        args.migrate_only, args.classify_media, args.classify_taxonomy,
        args.build_implementation_families,
        args.import_wlink, args.extract_titles,
        args.fetch_abstracts is not None, args.extract_abstracts,
        args.screen_venue,
    ))
    if not actions:
        parser.print_help()
        return

    conn = db_connect()
    try:
        count = ensure_schema(conn)
        print(f"Schema: {count} idempotent statements applied")
        if args.migrate_only:
            return
        if args.import_wlink:
            path = args.import_wlink
            if not path.is_absolute():
                path = ROOT / path
            print("WLink:", import_wlink(conn, path, args.match_threshold))
        if args.extract_titles:
            print("Titles:", extract_titles(conn, screened_only=args.screened_only))
        if args.fetch_abstracts is not None:
            print(
                "Abstract fetch:",
                fetch_abstracts(
                    conn, args.fetch_abstracts, args.batch_size, args.delay,
                    venue=args.venue, scope=args.abstract_scope,
                    year_from=args.year_from, year_to=args.year_to,
                    retry_terminal=args.retry_terminal,
                ),
            )
        if args.extract_abstracts:
            print(
                "Abstract extraction:",
                extract_abstracts(
                    conn, venue=args.extract_venue,
                    screened_only=args.screened_only,
                    year_from=args.year_from, year_to=args.year_to,
                ),
            )
        if args.screen_venue:
            print(
                "Venue screening:",
                screen_venue(
                    conn, args.screen_venue, args.screen_scope,
                    year_from=args.year_from, year_to=args.year_to,
                ),
            )
        rebuild_classification = any((
            args.classify_media, args.classify_taxonomy,
            args.fetch_abstracts is not None, args.screen_venue,
        ))
        if rebuild_classification:
            print(
                "Link media:",
                classify_link_media(conn, included_only=args.media_included_only),
            )
            print("Energy scopes:", classify_measurement_scopes(conn))
            print(
                "Link subtypes:",
                classify_link_subtypes(conn, included_only=args.media_included_only),
            )
        if args.build_implementation_families:
            print(
                "Implementation families:",
                build_implementation_families(conn, dry_run=args.family_dry_run),
            )
    finally:
        conn.close()


if __name__ == "__main__":
    main()
